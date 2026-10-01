#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
投标方案 Word 文档生成脚本

将 Markdown 格式的投标方案转换为排版规范的 Word 文档。
支持：标题层级、目录、表格、列表、页眉页脚、封面页。

依赖：python-docx
安装：pip install python-docx

用法：
    python generate_proposal.py --input <方案Markdown文件> --output <输出Word文件路径> --project-name <项目名称>

示例：
    python generate_proposal.py --input proposal.md --output "济南大学洗浴服务投标方案.docx" --project-name "济南大学学生公寓服务保障（洗浴）供应商项目"
"""

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Cm, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml
except ImportError:
    print("错误：未安装 python-docx 库。请执行：pip install python-docx")
    sys.exit(1)


# ============================================================
# 样式常量
# ============================================================
FONT_NAME_CN = "仿宋"
FONT_NAME_EN = "Times New Roman"
FONT_SIZE_BODY = Pt(14)       # 四号字 ≈ 14pt
FONT_SIZE_H1 = Pt(22)         # 小二号 ≈ 18pt → 用22pt更醒目
FONT_SIZE_H2 = Pt(16)         # 三号 ≈ 16pt
FONT_SIZE_H3 = Pt(14)         # 四号 ≈ 14pt
FONT_SIZE_COVER_TITLE = Pt(26)
FONT_SIZE_COVER_SUB = Pt(18)
LINE_SPACING = 1.5            # 1.5倍行距
MARGIN_TOP = Cm(2.54)
MARGIN_BOTTOM = Cm(2.54)
MARGIN_LEFT = Cm(3.17)
MARGIN_RIGHT = Cm(3.17)


def set_cell_shading(cell, color_hex):
    """设置表格单元格底色"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def setup_styles(doc):
    """配置文档样式"""
    # 设置默认字体
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME_EN
    style.font.size = FONT_SIZE_BODY
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
    pf = style.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)

    # 标题1样式
    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT_NAME_EN
    h1.font.size = FONT_SIZE_H1
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
    h1_fmt = h1.paragraph_format
    h1_fmt.space_before = Pt(24)
    h1_fmt.space_after = Pt(12)
    h1_fmt.line_spacing = 1.5

    # 标题2样式
    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT_NAME_EN
    h2.font.size = FONT_SIZE_H2
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
    h2_fmt = h2.paragraph_format
    h2_fmt.space_before = Pt(18)
    h2_fmt.space_after = Pt(8)
    h2_fmt.line_spacing = 1.5

    # 标题3样式
    h3 = doc.styles["Heading 3"]
    h3.font.name = FONT_NAME_EN
    h3.font.size = FONT_SIZE_H3
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
    h3_fmt = h3.paragraph_format
    h3_fmt.space_before = Pt(12)
    h3_fmt.space_after = Pt(6)
    h3_fmt.line_spacing = 1.5


def setup_page(doc):
    """设置页面格式"""
    for section in doc.sections:
        section.top_margin = MARGIN_TOP
        section.bottom_margin = MARGIN_BOTTOM
        section.left_margin = MARGIN_LEFT
        section.right_margin = MARGIN_RIGHT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)


def add_cover(doc, project_name, bidder_name=""):
    """添加封面页"""
    # 多次换行到页面中部
    for _ in range(6):
        doc.add_paragraph()

    # 项目名称
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(project_name)
    run.font.name = FONT_NAME_EN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
    run.font.size = FONT_SIZE_COVER_TITLE
    run.font.bold = True

    doc.add_paragraph()

    # 投标方案标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("投 标 方 案")
    run.font.name = FONT_NAME_EN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
    run.font.size = FONT_SIZE_COVER_SUB
    run.font.bold = True

    # 空行
    for _ in range(4):
        doc.add_paragraph()

    # 投标单位
    if bidder_name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"投标单位：{bidder_name}")
        run.font.name = FONT_NAME_EN
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
        run.font.size = Pt(16)

    # 日期
    import datetime
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(datetime.date.today().strftime("%Y年%m月%d日"))
    run.font.name = FONT_NAME_EN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
    run.font.size = Pt(16)

    # 封面后分页
    doc.add_page_break()


def add_toc(doc):
    """添加目录页"""
    p = doc.add_paragraph("目  录", style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 插入TOC域代码
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run._r.append(fldChar1)
    run2 = paragraph.add_run()
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>')
    run2._r.append(instrText)
    run3 = paragraph.add_run()
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    run3._r.append(fldChar2)
    run4 = paragraph.add_run("\uff08\u8bf7\u5728Word\u4e2d\u53f3\u952e\u70b9\u51fb\u6b64\u5904\uff0c\u9009\u62e9\u201c\u66f4\u65b0\u57df\u201d\u4ee5\u751f\u6210\u76ee\u5f55\uff09")
    run4.font.color.rgb = RGBColor(128, 128, 128)
    run5 = paragraph.add_run()
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run5._r.append(fldChar3)

    doc.add_page_break()


def add_header_footer(doc, project_name):
    """添加页眉页脚"""
    for section in doc.sections:
        # 页眉
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hp.add_run(project_name)
        run.font.name = FONT_NAME_EN
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(128, 128, 128)

        # 页眉下划线
        pPr = hp._element.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="1" w:color="auto"/>'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)

        # 页脚（页码）
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
        run._r.append(fldChar1)
        run2 = fp.add_run()
        instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
        run2._r.append(instrText)
        run3 = fp.add_run()
        fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
        run3._r.append(fldChar2)


def parse_table(lines, start_idx):
    """解析Markdown表格，返回(行列表, 结束行索引)"""
    rows = []
    idx = start_idx
    while idx < len(lines):
        line = lines[idx].strip()
        if not line.startswith("|"):
            break
        cells = [c.strip() for c in line.split("|")[1:-1]]  # 去掉首尾空元素
        # 跳过分隔行（---|---|---）
        if all(re.match(r"^[-:]+$", c) for c in cells):
            idx += 1
            continue
        rows.append(cells)
        idx += 1
    return rows, idx


def add_table(doc, rows):
    """将解析的表格行添加到Word文档"""
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = table.rows[i].cells[j]
                cell.text = ""
                p = cell.paragraphs[0]
                run = p.add_run(cell_text)
                run.font.name = FONT_NAME_EN
                run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
                run.font.size = Pt(12)
                # 表头行加粗 + 底色
                if i == 0:
                    run.font.bold = True
                    set_cell_shading(cell, "D9E2F3")


def parse_markdown(md_text):
    """解析Markdown文本，返回结构化内容列表"""
    elements = []
    lines = md_text.split("\n")
    idx = 0
    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()

        # 标题
        m = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            elements.append(("heading", level, text))
            idx += 1
            continue

        # 表格
        if stripped.startswith("|"):
            rows, new_idx = parse_table(lines, idx)
            if rows:
                elements.append(("table", rows))
            idx = new_idx
            continue

        # 有序列表
        m = re.match(r"^(\d+)[.、)）]\s*(.+)$", stripped)
        if m:
            elements.append(("ol_item", m.group(2).strip()))
            idx += 1
            continue

        # 无序列表
        if stripped.startswith(("- ", "* ", "• ")):
            elements.append(("ul_item", stripped[2:].strip()))
            idx += 1
            continue

        # 普通段落
        if stripped:
            elements.append(("paragraph", stripped))
            idx += 1
            continue

        # 空行
        idx += 1

    return elements


def render_to_doc(doc, elements):
    """将解析后的元素渲染到Word文档"""
    for elem in elements:
        if elem[0] == "heading":
            level = elem[1]
            text = elem[2]
            # 移除Markdown格式标记
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)  # 粗体
            text = re.sub(r"\*(.+?)\*", r"\1", text)        # 斜体
            heading_style = f"Heading {min(level, 3)}"
            doc.add_heading(text, level=min(level, 3))

        elif elem[0] == "table":
            add_table(doc, elem[1])

        elif elem[0] == "ol_item":
            p = doc.add_paragraph(style="List Number")
            text = elem[1]
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            text = re.sub(r"\*(.+?)\*", r"\1", text)
            run = p.add_run(text)
            run.font.name = FONT_NAME_EN
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
            run.font.size = FONT_SIZE_BODY

        elif elem[0] == "ul_item":
            p = doc.add_paragraph(style="List Bullet")
            text = elem[1]
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            text = re.sub(r"\*(.+?)\*", r"\1", text)
            run = p.add_run(text)
            run.font.name = FONT_NAME_EN
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
            run.font.size = FONT_SIZE_BODY

        elif elem[0] == "paragraph":
            text = elem[1]
            # 处理粗体和行内标记（简化处理）
            p = doc.add_paragraph()
            # 简单的粗体处理
            parts = re.split(r"(\*\*.+?\*\*)", text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    run = p.add_run(part)
                run.font.name = FONT_NAME_EN
                run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME_CN)
                run.font.size = FONT_SIZE_BODY


def generate_proposal(input_path, output_path, project_name, bidder_name=""):
    """主函数：将Markdown方案转为Word文档"""
    # 读取Markdown文件
    md_path = Path(input_path)
    if not md_path.exists():
        print(f"错误：文件不存在 - {input_path}")
        sys.exit(1)

    md_text = md_path.read_text(encoding="utf-8")

    # 创建Word文档
    doc = Document()

    # 设置样式和页面
    setup_styles(doc)
    setup_page(doc)

    # 添加封面
    add_cover(doc, project_name, bidder_name)

    # 添加目录
    add_toc(doc)

    # 添加页眉页脚
    add_header_footer(doc, project_name)

    # 解析Markdown并渲染
    elements = parse_markdown(md_text)
    render_to_doc(doc, elements)

    # 保存文档
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"✅ 投标方案已生成：{output_path}")


def main():
    parser = argparse.ArgumentParser(description="投标方案 Markdown → Word 转换工具")
    parser.add_argument("--input", required=True, help="输入的Markdown文件路径")
    parser.add_argument("--output", required=True, help="输出的Word文件路径")
    parser.add_argument("--project-name", required=True, help="项目名称（用于封面和页眉）")
    parser.add_argument("--bidder-name", default="", help="投标单位名称（用于封面）")
    args = parser.parse_args()

    generate_proposal(args.input, args.output, args.project_name, args.bidder_name)


if __name__ == "__main__":
    main()
