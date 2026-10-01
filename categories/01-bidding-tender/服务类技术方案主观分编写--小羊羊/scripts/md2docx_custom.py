#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
服务类技术方案 — Markdown to DOCX 转换器 (v2.0.0 升级版)
基于货物类技能 md2docx.py v2.35.0 的 38+ 格式化规则升级。

升级内容（v2.0.0）：
- 标题使用 Word 内置 Heading 样式（TOC/导航窗格可用）
- 宋体百分号标准化（半角% → 全角％）
- ×占位符红色标记（自动标红加粗）
- 表格标题行合并（单格首行跨列合并 + 浅蓝底）
- 表格垂直居中修复（先删旧 vAlign 再设新值）
- 行距全覆盖（正文/标题/表格/TOC/空行/图片标题）
- 分页符三级覆盖（H1/H2/H3 前加分页符）
- 内容俗套语清理
- 表格后验证（MD表格数 vs DOCX表格数）

Usage:
  python md2docx_custom.py <input_md> <output_docx> [--project-name "项目名称"]
"""
import re, os, sys, argparse
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement


# ====== HELPER FUNCTIONS ======

def set_cell_shading(cell, color):
    """设置表格单元格底色"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_vertical_center(cell):
    """设置单元格垂直居中（先删旧值再设新值，防止 Table Grid 默认 top 覆盖）"""
    tcPr = cell._element.get_or_add_tcPr()
    for old_v in tcPr.findall(qn('w:vAlign')):
        tcPr.remove(old_v)
    vAlign = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>')
    tcPr.append(vAlign)


def fix_heading_style(doc, style_name, font_size, bold):
    """修改内置标题样式字体为宋体"""
    style = doc.styles[style_name]
    style.font.name = '宋体'
    style.font.size = Pt(font_size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.5
    rPr = style.element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), '宋体')
    rFonts.set(qn('w:ascii'), '宋体')
    rPr.insert(0, rFonts)


def add_heading_with_style(doc, text, level):
    """添加标题（使用内置样式 + run 级宋体覆写）"""
    p = doc.add_heading(text, level=level)
    p.paragraph_format.line_spacing = 1.5
    for run in p.runs:
        run.font.name = '宋体'
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:eastAsia'), '宋体')
        run._element.get_or_add_rPr().insert(0, rFonts)
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_body_paragraph(doc, text, first_line_indent=True, bold=False):
    """添加正文段落（支持 × 占位符自动标红）"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)

    # 以 × 为界拆分 runs，实现红色标记
    parts = re.split(r'(×+)', text)
    for part in parts:
        if part == '':
            continue
        run = p.add_run(part)
        run.font.name = '宋体'
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:eastAsia'), '宋体')
        run._element.get_or_add_rPr().insert(0, rFonts)
        run.font.size = Pt(12)
        run.font.bold = bold

        if '×' in part:
            # XML 直写红色，清除 themeColor 防 Word 主题覆盖
            rPr = run._element.get_or_add_rPr()
            for old_color in rPr.findall(qn('w:color')):
                rPr.remove(old_color)
            color_el = OxmlElement('w:color')
            color_el.set(qn('w:val'), 'FF0000')
            rPr.append(color_el)
            run.font.bold = True
        else:
            rPr = run._element.get_or_add_rPr()
            for old_color in rPr.findall(qn('w:color')):
                rPr.remove(old_color)
            color_el = OxmlElement('w:color')
            color_el.set(qn('w:val'), '000000')
            rPr.append(color_el)
    return p


def clean_markdown(text):
    """移除 Markdown 格式标记（保留纯文本）"""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text


def clean_cliches(text):
    """清除俗套语（对采购人的实用价值等）"""
    text = re.sub(r'\*\*?【对采购人的实用价值】[：:]\s*\*\*?', '', text)
    text = re.sub(r'【对采购人的实用价值】[：:]?\s*', '', text)
    text = text.replace('对采购人的价值', '')
    return text


# ====== DOCUMENT SETUP ======

def setup_document():
    """创建并配置 Word 文档"""
    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    # 修改正文样式
    normal = doc.styles['Normal']
    normal.font.name = '宋体'
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    rPr = normal.element.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), '宋体')
    rPr.insert(0, rFonts)

    # 修改标题样式
    fix_heading_style(doc, 'Heading 1', 14, True)
    fix_heading_style(doc, 'Heading 2', 14, False)
    fix_heading_style(doc, 'Heading 3', 12, False)

    return doc


def add_toc(doc):
    """添加 TOC 目录域"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run('目  录')
    run.font.name = '宋体'
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), '宋体')
    run._element.get_or_add_rPr().insert(0, rFonts)
    run.font.size = Pt(16)
    run.font.bold = True

    toc_p = doc.add_paragraph()
    toc_p.paragraph_format.line_spacing = 1.5

    # fldChar begin
    run1 = toc_p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    run1._r.append(fld_begin)

    # instrText
    run2 = toc_p.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    run2._r.append(instr)

    # fldChar separate
    run3 = toc_p.add_run()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    run3._r.append(fld_sep)

    # 提示语
    run4 = toc_p.add_run('（请在Word中右键点击此处，选择"更新域"以生成目录）')
    run4.font.size = Pt(10.5)
    run4.font.color.rgb = RGBColor(128, 128, 128)

    # fldChar end
    run5 = toc_p.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run5._r.append(fld_end)

    doc.add_page_break()


def add_page_numbers(doc):
    """添加页码（底部居中）"""
    for sec in doc.sections:
        footer = sec.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.line_spacing = 1.5

        for tag_type in ['begin', 'instr', 'end']:
            run = fp.add_run()
            if tag_type == 'instr':
                el = OxmlElement('w:instrText')
                el.set(qn('xml:space'), 'preserve')
                el.text = ' PAGE '
            else:
                el = OxmlElement('w:fldChar')
                el.set(qn('w:fldCharType'), tag_type)
            run._r.append(el)
            run.font.name = '宋体'
            run.font.size = Pt(9)


# ====== TABLE HANDLING ======

def add_table_to_doc(doc, data_rows, has_title_row=False):
    """
    添加格式化表格到文档。
    has_title_row: 首行是否为标题行（单格 → 跨列合并 + 浅蓝底）
    """
    if not data_rows:
        return

    # 如果有标题行，先提取
    title_text = None
    body_rows = list(data_rows)
    if has_title_row and len(body_rows) > 0 and len(body_rows[0]) == 1:
        title_text = clean_markdown(body_rows[0][0])
        body_rows = body_rows[1:]

    num_cols = max(len(r) for r in body_rows) if body_rows else 2
    num_rows = len(body_rows) + (1 if title_text else 0)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # 表格宽度 100%
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)

    row_offset = 0

    # 标题行（浅蓝底 + 跨列合并）
    if title_text:
        cell = table.cell(0, 0)
        if num_cols > 1:
            cell.merge(table.cell(0, num_cols - 1))
        cell.text = ''
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.line_spacing = 1.5
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        cr = cp.add_run(title_text)
        cr.font.name = '宋体'
        cr.font.size = Pt(10.5)
        cr.font.bold = True
        set_cell_shading(cell, 'D9E2F3')
        set_cell_vertical_center(cell)
        row_offset = 1

    # 数据行
    for row_idx, row_data in enumerate(body_rows):
        for col_idx in range(num_cols):
            cell = table.cell(row_idx + row_offset, col_idx)
            cell_text = clean_markdown(row_data[col_idx]) if col_idx < len(row_data) else ''
            cell.text = ''

            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.line_spacing = 1.5
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.first_line_indent = Pt(0)

            cr = cp.add_run(cell_text)
            cr.font.name = '宋体'
            cr.font.size = Pt(10.5)

            # 垂直居中（先删后设）
            set_cell_vertical_center(cell)

            # 表头行样式
            if row_idx == 0 and not title_text:
                cr.font.bold = True
                cr.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_shading(cell, '4472C4')
            elif row_idx == 0 and title_text:
                # 有标题行时，第一行数据行也可能是表头
                cr.font.bold = True
                cr.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_shading(cell, '4472C4')

    return table


# ====== IMAGE HANDLING ======

def add_image_from_ref(doc, img_path, caption_text):
    """真实嵌入图片"""
    if not os.path.isfile(img_path):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(f'[图片未找到: {caption_text}]')
        run.font.color.rgb = RGBColor(255, 0, 0)
        return

    try:
        pic = doc.add_picture(img_path, width=Cm(14.0))
        last_p = doc.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        last_p.paragraph_format.line_spacing = 1.5
    except Exception as e:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(f'[图片加载失败: {caption_text}]')
        run.font.color.rgb = RGBColor(255, 0, 0)
        return

    # 图片标题
    if caption_text:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(caption_text)
        run.font.name = '宋体'
        run.font.size = Pt(9)


# ====== MAIN CONVERSION ======

def convert_md_to_docx(md_path, output_path, project_name=None):
    """主转换函数"""
    doc = setup_document()

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # 宋体百分号标准化：半角% → 全角％
    content = content.replace('%', '\uff05')

    # 俗套语清理
    content = clean_cliches(content)

    lines = content.split('\n')

    # ---- 封面 ----
    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(project_name or '')
    run.font.name = '宋体'
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), '宋体')
    run._element.get_or_add_rPr().insert(0, rFonts)
    run.font.size = Pt(22)
    run.font.bold = True

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('技术方案')
    run.font.name = '宋体'
    run.font.size = Pt(18)

    doc.add_page_break()

    # ---- TOC ----
    add_toc(doc)

    # ---- Body parsing ----
    in_table = False
    table_lines = []
    image_refs = {}  # {line_index: (path, caption)}
    just_after_h1 = False
    just_after_h2 = False
    first_h1 = True

    # 第一轮：扫描图片引用
    for i, line in enumerate(lines):
        img_match = re.match(r'^!\[(.+?)\]\((.+)\)', line.strip())
        if img_match:
            caption = img_match.group(1).strip()
            path = img_match.group(2).strip()
            image_refs[i] = (path, caption)

    def is_table_line(text):
        t = text.strip()
        return t.startswith('|') and t.endswith('|') and t.count('|') >= 2

    def flush_table():
        nonlocal in_table, table_lines
        if in_table and table_lines:
            in_table = False
            if len(table_lines) >= 2:
                rows = []
                for tl in table_lines:
                    cells = [c.strip() for c in tl.strip().strip('|').split('|')]
                    rows.append(cells)
                # 过滤分隔行
                data_rows = [r for r in rows if not all(re.match(r'^[-:\s]+$', c.strip()) for c in r)]
                if data_rows:
                    # 检测首行是否为单格标题行
                    has_title = (len(data_rows[0]) == 1)
                    add_table_to_doc(doc, data_rows, has_title_row=has_title)
            table_lines = []

    i = 0
    md_table_count = 0

    while i < len(lines):
        line = lines[i].rstrip()

        # 图片引用行
        if i in image_refs:
            img_path, caption = image_refs[i]
            add_image_from_ref(doc, img_path, caption)
            i += 1
            continue

        # 空行
        if not line.strip():
            if in_table:
                flush_table()
            i += 1
            continue

        # 分隔线
        if line.strip() == '---':
            if in_table:
                flush_table()
            i += 1
            continue

        # 撞到非表格行 → flush table
        if in_table and not is_table_line(line):
            flush_table()

        # ## Heading 2
        if line.startswith('## ') and not line.startswith('### '):
            text = line.lstrip('#').strip()
            if not just_after_h1:
                doc.add_page_break()
            just_after_h1 = False
            just_after_h2 = True
            add_heading_with_style(doc, text, 2)
            i += 1
            continue

        # ### Heading 3
        if line.startswith('### ') and not line.startswith('#### '):
            if in_table:
                flush_table()
            text = line.lstrip('#').strip()
            if not just_after_h2:
                doc.add_page_break()
            just_after_h2 = False
            add_heading_with_style(doc, text, 3)
            i += 1
            continue

        # #### Heading 4 (treat as bold body)
        if line.startswith('#### '):
            if in_table:
                flush_table()
            text = line.lstrip('#').strip()
            add_body_paragraph(doc, text, bold=True)
            i += 1
            continue

        # 代码块
        if line.strip().startswith('```'):
            if in_table:
                flush_table()
            i += 1
            continue

        # 表格
        if is_table_line(line):
            if not in_table:
                in_table = True
                table_lines = []
                md_table_count += 1
            table_lines.append(line)
            i += 1
            continue

        # 加粗行
        if line.strip().startswith('**') and line.strip().endswith('**'):
            text = line.strip().strip('*')
            add_body_paragraph(doc, text, bold=True)
            just_after_h1 = False
            just_after_h2 = False
            i += 1
            continue

        # 图片引用（备用检测）
        if line.strip().startswith('!['):
            img_match = re.match(r'!\[(.+?)\]\((.+)\)', line.strip())
            if img_match:
                add_image_from_ref(doc, img_match.group(2), img_match.group(1))
            i += 1
            continue

        # # Heading 1 (chapter title)
        if line.startswith('# ') and not line.startswith('## '):
            text = line.lstrip('#').strip()
            if not first_h1:
                doc.add_page_break()
            first_h1 = False
            just_after_h1 = True
            add_heading_with_style(doc, text, 1)
            i += 1
            continue

        # 普通段落
        text = clean_markdown(line.strip())
        if text:
            add_body_paragraph(doc, text)
            just_after_h1 = False
            just_after_h2 = False
        i += 1

    # 清空末尾表格
    if in_table:
        flush_table()

    # ---- 页码 ----
    add_page_numbers(doc)

    # ---- 保存 ----
    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    doc.save(output_path)

    # ---- 统计 ----
    headings = [p for p in doc.paragraphs if p.style.name.startswith('Heading')]
    h1 = sum(1 for p in headings if p.style.name == 'Heading 1')
    h2 = sum(1 for p in headings if p.style.name == 'Heading 2')
    h3 = sum(1 for p in headings if p.style.name == 'Heading 3')

    print(f'✅ DOCX: {output_path}')
    print(f'   H1={h1} H2={h2} H3={h3} tables={len(doc.tables)} images={len(image_refs)}')
    print(f'   MD表格={md_table_count} DOCX表格={len(doc.tables)} {"✅" if md_table_count == len(doc.tables) else "❌ 数量不匹配！"}')

    # 表格数量验证
    if md_table_count != len(doc.tables):
        print(f'⚠️ 警告: MD表格数({md_table_count}) ≠ DOCX表格数({len(doc.tables)})，请检查转换结果。')

    # Heading 样式验证
    for p in doc.paragraphs:
        if p.style.name == 'Normal' and any(
            run.font.size and run.font.size >= Pt(14) for run in p.runs
        ):
            print(f'⚠️ 警告: 正文段落可能含大号字体（疑似标题未正确映射）')

    return doc


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='服务类技术方案 MD→DOCX 转换器 v2.0.0')
    parser.add_argument('input_md', help='输入 Markdown 文件路径')
    parser.add_argument('output_docx', help='输出 DOCX 文件路径')
    parser.add_argument('--project-name', help='项目名称（用于封面）', default='')
    args = parser.parse_args()

    convert_md_to_docx(args.input_md, args.output_docx, args.project_name)
