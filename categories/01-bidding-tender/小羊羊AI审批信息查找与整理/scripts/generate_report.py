# -*- coding: utf-8 -*-
"""
政府项目审批报告生成脚本
用法: python generate_report.py --projects "<JSON数据>" --output "<输出路径>" --days <天数>

参数说明:
  --days <天数>  指定搜索时间范围，默认365天（1年），仅收录该天数内的项目
"""
import json
import sys
from datetime import datetime, timedelta

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("❌ 需要安装openpyxl库: pip install openpyxl")
    sys.exit(1)


def create_excel_report(projects, output_path):
    """生成Excel报告"""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "项目审批汇总"

    # 样式定义
    header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    cell_font = Font(name='微软雅黑', size=10)
    cell_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    even_row_fill = PatternFill(start_color='D6E3F8', end_color='D6E3F8', fill_type='solid')
    odd_row_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

    approved_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    pending_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 表头
    headers = ['序号', '学校/单位', '项目名称', '项目类型', '审批状态',
               '发布时间', '官方网址', '数据来源', '备注']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 数据行
    for row_idx, project in enumerate(projects, 2):
        ws.cell(row=row_idx, column=1, value=row_idx-1).alignment = cell_alignment
        ws.cell(row=row_idx, column=2, value=project.get("学校", "")).alignment = cell_alignment
        ws.cell(row=row_idx, column=3, value=project.get("项目名称", "")).alignment = cell_alignment
        ws.cell(row=row_idx, column=4, value=project.get("项目类型", "")).alignment = cell_alignment
        ws.cell(row=row_idx, column=5, value=project.get("审批状态", "")).alignment = cell_alignment
        ws.cell(row=row_idx, column=6, value=project.get("发布时间", "")).alignment = cell_alignment
        ws.cell(row=row_idx, column=7, value=project.get("官方网址", "")).alignment = cell_alignment
        ws.cell(row=row_idx, column=8, value=project.get("数据来源", "")).alignment = cell_alignment
        ws.cell(row=row_idx, column=9, value=project.get("备注", "")).alignment = cell_alignment

        # 应用样式
        for col in range(1, 10):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = cell_font
            cell.border = thin_border
            cell.fill = even_row_fill if row_idx % 2 == 0 else odd_row_fill

        # 备注颜色
        note_cell = ws.cell(row=row_idx, column=9)
        if '✅' in project.get("备注", ""):
            note_cell.fill = approved_fill
        elif '⚠️' in project.get("备注", ""):
            note_cell.fill = pending_fill

        # URL样式
        url_cell = ws.cell(row=row_idx, column=7)
        url_cell.font = Font(name='微软雅黑', size=10, color='0563C1', underline='single')

    # 列宽
    column_widths = [6, 20, 32, 12, 28, 14, 55, 25, 30]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # 行高
    ws.row_dimensions[1].height = 30
    for row_idx in range(2, len(projects) + 2):
        ws.row_dimensions[row_idx].height = 35

    # 冻结首行
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:I{len(projects) + 1}"

    # 统计Sheet
    stats_ws = wb.create_sheet(title="统计说明")
    verified_count = sum(1 for p in projects if '✅' in p.get("备注", ""))
    pending_count = sum(1 for p in projects if '⚠️' in p.get("备注", ""))

    stats_data = [
        ["📊 数据统计", ""],
        ["总项目数", len(projects)],
        ["✅ 已验证", verified_count],
        ["⚠️ 待核实", pending_count],
        ["", ""],
        ["📅 整理时间", datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
    ]

    for row_idx, (label, value) in enumerate(stats_data, 1):
        stats_ws.cell(row=row_idx, column=1, value=label).font = Font(name='微软雅黑', size=11, bold=True)
        stats_ws.cell(row=row_idx, column=2, value=value).font = Font(name='微软雅黑', size=11)

    stats_ws.column_dimensions['A'].width = 20
    stats_ws.column_dimensions['B'].width = 30

    wb.save(output_path)
    print(f"✅ 报告已生成: {output_path}")
    return output_path


def filter_projects_by_days(projects, days=365):
    """过滤保留近N天内的项目"""
    cutoff_date = datetime.now() - timedelta(days=days)
    filtered = []
    for p in projects:
        pub_date_str = p.get("发布时间", "")
        if pub_date_str:
            try:
                pub_date = datetime.strptime(pub_date_str, "%Y-%m-%d")
                if pub_date >= cutoff_date:
                    filtered.append(p)
            except ValueError:
                # 日期格式解析失败，保留该项目
                filtered.append(p)
        else:
            # 没有日期的项目默认保留
            filtered.append(p)
    return filtered


def main():
    """主函数"""
    # 默认过滤天数：365天（1年）
    filter_days = 365

    if len(sys.argv) < 2:
        # 默认测试数据
        projects = [
            {
                "学校": "测试学校",
                "项目名称": "测试项目",
                "项目类型": "教学楼",
                "审批状态": "已批复",
                "发布时间": "2025-01-01",
                "官方网址": "https://example.com",
                "数据来源": "测试",
                "备注": "✅ 已验证"
            }
        ]
        output_path = "项目审批报告.xlsx"
    else:
        # 解析命令行参数
        args = sys.argv[1:]
        projects = []
        output_path = "项目审批报告.xlsx"

        i = 0
        while i < len(args):
            if args[i] == "--projects" and i + 1 < len(args):
                try:
                    projects = json.loads(args[i + 1])
                    i += 2
                except json.JSONDecodeError:
                    print("❌ JSON格式错误")
                    sys.exit(1)
            elif args[i] == "--output" and i + 1 < len(args):
                output_path = args[i + 1]
                i += 2
            elif args[i] == "--days" and i + 1 < len(args):
                try:
                    filter_days = int(args[i + 1])
                    i += 2
                except ValueError:
                    print("❌ --days 参数必须是整数")
                    sys.exit(1)
            else:
                i += 1

    if not projects:
        print("❌ 没有项目数据")
        sys.exit(1)

    # 时间过滤：仅保留近N天内的项目
    original_count = len(projects)
    projects = filter_projects_by_days(projects, days=filter_days)
    filtered_count = original_count - len(projects)

    if filtered_count > 0:
        print(f"📅 已过滤 {filtered_count} 个超过 {filter_days} 天的项目，保留 {len(projects)} 个")

    if not projects:
        print("❌ 没有符合时间范围的项目数据")
        sys.exit(1)

    create_excel_report(projects, output_path)


if __name__ == "__main__":
    main()
