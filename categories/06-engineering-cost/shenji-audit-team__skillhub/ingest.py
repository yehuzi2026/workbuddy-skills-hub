# -*- coding: utf-8 -*-
"""文件原生解析（衔接环节自动抽取，免去用户手抄数据）。

  - ingest_xlsx(path, sheet=None)  读取 Excel（清单/定额/结算表）为结构化行
  - ingest_csv(path)               读取 CSV（造价站信息价常见格式，自动识别 UTF-8/GBK）
  - ingest_pdf(path)               提取 PDF 文本行（需 pdfplumber）
  - build_rates(path, ...)         把当地信息价/费用定额（xlsx 或 csv）转成 local_rates.json
  - file_fingerprint(path)          算电子数据导出文件的 SHA-256 / MD5，供「电子数据取数说明」登记
  - 命令行：
      python ingest.py <文件>                                  解析清单/结算/PDF/CSV
      python ingest.py --hash <导出文件>                        登记文件指纹（SHA-256 / MD5）
      python ingest.py --hash <导出文件> --expect <登记值>        使用前重算比对（不一致即退出码 1）
      python ingest.py --build-rates <信息价表.xlsx|csv> --kind materials
          [--region 贵州毕节] [--period 2026第3期] [--source 毕节造价站] [--out ...]

依赖缺失时明确报错并给出安装命令，不静默失败。
"""
import sys, os, json, csv, argparse, datetime, hashlib


def _setup_io():
    for _name in ("stdin", "stdout", "stderr"):
        try:
            getattr(sys, _name).reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


_setup_io()


def _cell(v):
    return "" if v is None else str(v).strip()


def _rows_to_records(rows):
    rows = list(rows)
    if not rows:
        return [], []
    header = [_cell(c) or "列%d" % (i + 1) for i, c in enumerate(rows[0])]
    recs = []
    for r in rows[1:]:
        if all(_cell(c) == "" for c in r):
            continue
        recs.append({header[i]: _cell(r[i]) for i in range(min(len(header), len(r)))})
    return header, recs


def ingest_xlsx(path, sheet=None, max_rows=200):
    try:
        from openpyxl import load_workbook
    except Exception:
        raise RuntimeError("读取 xlsx 需要 openpyxl：pip install openpyxl")
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > max_rows:
            break
        rows.append(row)
    header, recs = _rows_to_records(rows)
    return {"type": "xlsx", "sheet": ws.title, "header": header, "rows": recs,
            "count": len(recs)}


def ingest_csv(path, max_rows=200, sheet=None):
    last_err = None
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            with open(path, encoding=enc, newline="") as f:
                rows = [r for i, r in enumerate(csv.reader(f)) if i <= max_rows]
            header, recs = _rows_to_records(rows)
            return {"type": "csv", "encoding": enc, "header": header, "rows": recs,
                    "count": len(recs)}
        except UnicodeDecodeError as e:
            last_err = e
            continue
    raise RuntimeError("CSV 编码无法识别（已尝试 utf-8-sig/utf-8/gb18030/gbk）：%s" % last_err)


def ingest_pdf(path, max_lines=300):
    try:
        import pdfplumber
    except Exception:
        raise RuntimeError("读取 PDF 需要 pdfplumber：pip install pdfplumber")
    out = []
    with pdfplumber.open(path) as pdf:
        for pi, page in enumerate(pdf.pages):
            if pi >= 50:
                break
            txt = page.extract_text() or ""
            for ln in txt.splitlines():
                if ln.strip():
                    out.append(ln.strip())
                if len(out) >= max_lines:
                    break
            if len(out) >= max_lines:
                break
    return {"type": "pdf", "lines": out, "count": len(out)}


def _read_price_records(path, sheet=None):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        rec = ingest_xlsx(path, sheet=sheet, max_rows=5000)
    elif ext in (".csv", ".txt"):
        rec = ingest_csv(path, sheet=sheet)
    elif ext == ".xls":
        raise RuntimeError("旧版 .xls 不受支持：请用 Excel 另存为 .xlsx 或 .csv 后重试")
    else:
        raise RuntimeError("仅支持 xlsx/xlsm/csv，收到 %s" % path)
    pairs = []
    for r in rec["rows"]:
        keys = list(r.keys())
        if not keys:
            continue
        name = _cell(r.get(keys[0]))
        val_raw = r.get(keys[1]) if len(keys) > 1 else ""
        if not name or not _cell(val_raw):
            continue
        try:
            price = float(str(val_raw).replace(",", "").replace("，", "").strip())
        except ValueError:
            continue
        pairs.append((name, price))
    if not pairs:
        raise ValueError("未解析到「名称 + 数值」两列数据，请确认表格前两列为 名称/单价（费率）")
    return rec, pairs


def file_fingerprint(path, expect=""):
    """登记电子数据导出文件的「文件指纹」：大小、修改时间、SHA-256 与 MD5。

    用于「电子数据取数说明」的校验值登记：**取数时算一次，此后每次使用前重算比对**；
    比对不一致说明文件已被改动，不得使用，须重新取数（失败闭合）。
    """
    if not os.path.isfile(path):
        raise RuntimeError("文件不存在或不是普通文件：%s" % path)
    h256, h5 = hashlib.sha256(), hashlib.md5()
    size = 0
    with open(path, "rb") as f:
        while True:
            chunk = f.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            h256.update(chunk)
            h5.update(chunk)
    sha256, md5 = h256.hexdigest(), h5.hexdigest()
    out = {
        "文件名": os.path.basename(path),
        "路径": os.path.abspath(path),
        "大小(字节)": size,
        "最后修改时间": datetime.datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d %H:%M:%S"),
        "SHA-256": sha256,
        "MD5": md5,
        "登记用途": "填入「审计项目档案/04_审计底稿与证据/电子数据取数说明.md」的「校验值（SHA-256 / MD5）」栏，"
                    "并与导出时间、记录数、合计金额一并登记",
        "使用前校验": "python ingest.py --hash <文件> --expect <原登记值>（SHA-256 或 MD5 任一一致即为一致）",
    }
    if expect:
        e = expect.strip().lower()
        same = (e == sha256.lower()) or (e == md5.lower())
        out["登记值"] = expect
        out["比对结果"] = "一致（文件未被改动）" if same else "不一致（文件已改动，不得使用，须重新取数）"
    return out


def build_rates(path, sheet=None, kind="materials", region="", period="", source="",
                is_example=False, out=None, price_unit=""):
    """把当地造价站发布的「信息价/费用定额」表格转为 local_rates.json。

    kind="materials"：前两列 = 名称/单价（元/计量单位，建议名称里带单位，如 "C30商品混凝土 元/m³"）
    kind="fees"     ：前两列 = 费率名/费率值（0.05=5%），按关键词映射到 fee_rates
    返回写入的 dict；out 为空时仅返回不落盘。
    """
    rec, pairs = _read_price_records(path, sheet)
    data = {}
    for name, price in pairs:
        data[name] = price

    if kind == "fees":
        fee_map = {"manage_rate": ["管理费率", "管理费"],
                   "profit_rate": ["利润率", "利润"],
                   "risk_amp": ["风险幅度", "风险"],
                   "tax_rate": ["税率", "增值税"],
                   "gov_fee_rate": ["规费费率", "规费"]}
        fee_rates, unmatched = {}, []
        for k, v in data.items():
            matched = None
            for fk, aliases in fee_map.items():
                if any(a in k for a in aliases):
                    matched = fk
                    break
            if matched:
                fee_rates[matched] = v
            else:
                unmatched.append(k)
        # 费率若以百分数填写（>1）自动折算为小数，避免 5 被当成 500%
        converted = []
        for fk, fv in list(fee_rates.items()):
            if fv > 1:
                fee_rates[fk] = fv / 100.0
                converted.append(fk)
        result = {"fee_rates": fee_rates}
        if unmatched:
            result["fee_rates_unmatched"] = unmatched
        if converted:
            result["fee_rates_converted_to_decimal"] = converted
    else:
        result = {"materials": data}

    result.update({"region": region or "（待填）地区",
                   "period": period or "（待填）期次",
                   "source": source or "（待填）发布机构",
                   "updated": datetime.date.today().isoformat(),
                   "is_example": is_example,
                   "price_unit": price_unit or "元/计量单位（键名中可含单位）",
                   "rows_parsed": len(pairs),
                   "source_file": os.path.abspath(path),
                   "source_sheet": rec.get("sheet", rec.get("encoding", ""))})
    result.setdefault("fee_rates", {"manage_rate": 0.05, "profit_rate": 0.04})
    result["_warning"] = ("⚠ 示例占位数据，禁止用于正式结论；请提供当地造价管理部门公布的真实数据。"
                          if is_example else
                          "本文件为引擎取价依据，请确保数据来自当地造价管理部门最新公布文件，并保留期次与来源。")
    if out:
        d = os.path.dirname(out)
        if d:
            os.makedirs(d, exist_ok=True)
        with open(out, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
    return result


def main():
    ap = argparse.ArgumentParser(description="文件原生解析（清单/结算/PDF/CSV/当地价格数据）")
    ap.add_argument("path", nargs="?", help="xlsx/xlsm/csv/pdf 文件路径")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--hash", default=None, metavar="文件",
                    help="登记文件指纹：算 SHA-256 与 MD5（用于电子数据取数说明的校验值登记）")
    ap.add_argument("--expect", default="", help="配合 --hash：填入原登记校验值，重算比对是否被改动")
    ap.add_argument("--max_rows", type=int, default=200)
    ap.add_argument("--build-rates", default=None,
                    help="把当地信息价/费用定额表转为 local_rates.json，填此参数=文件路径")
    ap.add_argument("--kind", choices=["materials", "fees"], default="materials")
    ap.add_argument("--region", default="")
    ap.add_argument("--period", default="")
    ap.add_argument("--source", default="")
    ap.add_argument("--price_unit", default="")
    ap.add_argument("--example", action="store_true", help="标记为示例占位数据")
    ap.add_argument("--out", default=None,
                    help="输出 local_rates.json 路径（默认包内 references/data/local_rates.json）")
    args = ap.parse_args()

    if args.hash:
        try:
            fp = file_fingerprint(args.hash, args.expect)
        except Exception as e:
            print(json.dumps({"错误": str(e)}, ensure_ascii=False))
            return 2
        print(json.dumps(fp, ensure_ascii=False, indent=2))
        if args.expect and "不一致" in fp.get("比对结果", ""):
            return 1
        return 0


    if args.build_rates:
        out = args.out or os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                       "references", "data", "local_rates.json")
        try:
            res = build_rates(args.build_rates, args.sheet, args.kind, args.region,
                              args.period, args.source, args.example, out, args.price_unit)
        except Exception as e:
            print(json.dumps({"错误": str(e)}, ensure_ascii=False))
            return 2
        n = len(res.get("materials") or res.get("fee_rates") or {})
        print(json.dumps({"已生成": out, "类型": args.kind, "条目数": n,
                          "地区": res["region"], "期次": res["period"],
                          "是否示例数据": res["is_example"],
                          "提示": "计算时用 --rates_file 指向该文件，引擎会自动引用并标注数据来源"},
                         ensure_ascii=False, indent=2))
        return 0

    p = args.path
    if not p:
        print(json.dumps({"错误": "请提供文件路径，或用 --build-rates 导入当地价格数据"},
                         ensure_ascii=False))
        return 1
    if not os.path.exists(p):
        print(json.dumps({"错误": "文件不存在：%s" % p}, ensure_ascii=False))
        return 2
    try:
        low = p.lower()
        if low.endswith((".xlsx", ".xlsm")):
            res = ingest_xlsx(p, args.sheet, args.max_rows)
        elif low.endswith((".csv", ".txt")):
            res = ingest_csv(p, args.max_rows)
        elif low.endswith(".pdf"):
            res = ingest_pdf(p)
        elif low.endswith(".xls"):
            print(json.dumps({"错误": "旧版 .xls 不受支持：请另存为 .xlsx 或 .csv"},
                             ensure_ascii=False))
            return 2
        else:
            print(json.dumps({"错误": "仅支持 xlsx/xlsm/csv/pdf，收到 %s" % p}, ensure_ascii=False))
            return 1
    except Exception as e:
        print(json.dumps({"错误": str(e)}, ensure_ascii=False))
        return 2
    print(json.dumps(res, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
