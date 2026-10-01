#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
docx_table_tools.py — docx 表格读写工具（保留原格式）

用法（库）:
    from docx_table_tools import load, iter_body, find_tables, read_table, set_cell
    d = load(path)
    iter_body(d)                              # 按文档顺序打印 段落/表格/图片
    ts = find_tables(d, ["合同金额（万元）", "合同甲方"])   # 按关键词找表
    print(read_table(ts[0]))
    set_cell(ts[0].rows[1].cells[5], "2.00")

用法（CLI 速查）:
    python docx_table_tools.py inspect  <docx>
    python docx_table_tools.py dump     <docx> [table_index]
"""
import sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph


def load(path):
    return Document(path)


def iter_body(doc, show_images=True):
    """按 body 真实顺序输出 (类型, 内容)，段落/表格/图片顺序不会乱。"""
    out = []
    ti = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            p = Paragraph(child, doc)
            has_img = bool(child.findall('.//' + qn('a:blip')))
            txt = p.text.strip()
            if txt or (show_images and has_img):
                out.append(("P", txt[:80] + ("  [图片]" if has_img else "")))
        elif child.tag == qn('w:tbl'):
            out.append(("T", "表格%d" % ti))
            ti += 1
    for t in out:
        print(" ", t[0], ":", t[1])
    return out


def find_tables(doc, keywords, mode="all"):
    """找同时包含全部 keywords（mode='all'）或任一（mode='any'）的表格。"""
    hits = []
    for tbl in doc.tables:
        blob = " ".join(c.text for row in tbl.rows for c in row.cells)
        ok = all(k in blob for k in keywords) if mode == "all" else any(k in blob for k in keywords)
        if ok:
            hits.append(tbl)
    return hits


def read_table(tbl, max_rows=99, cell_len=40):
    lines = []
    for ri, row in enumerate(tbl.rows[:max_rows]):
        lines.append("R%d: " % ri + " | ".join(c.text.strip().replace("\n", "/")[:cell_len] for c in row.cells))
    return "\n".join(lines)


def set_cell(cell, text, font='宋体', size_pt=9, align=WD_ALIGN_PARAGRAPH.CENTER):
    """填单元格并【保留格式】。切勿用 cell.text = ... （会清掉字体/对齐）。"""
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    if p.alignment is None:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = font
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)   # 中文字体必须设 eastAsia
    run.font.size = Pt(size_pt)
    return run


def copy_row(src_cells, tgt_cells, mapping):
    """按 mapping 复制：mapping[i] = 源列 index（None 表示目标列留空）。
    例：源 9 列 → 目标 10 列（第9列页码留空、第10列放源第9列备注）:
        copy_row(src, tgt, [0,1,2,3,4,5,6,7,None,8])
    """
    for j, src_i in enumerate(mapping):
        if j >= len(tgt_cells):
            break
        set_cell(tgt_cells[j], "" if src_i is None else src_cells[src_i].text.strip())


def move_table_after(doc, tbl_or_index, anchor_text):
    """把整张表搬到『文本 == anchor_text』的段落正下方，保留表格全部格式/边框/底纹。
    ⚠️ 若目标位置本就紧邻另一张表，Word 打开时会【自动把两张表合并】成一张。
       搬表前务必确认该处没有别的表；若必须并存，先在两表之间插一个空段落隔离。
    """
    tbl = doc.tables[tbl_or_index] if isinstance(tbl_or_index, int) else tbl_or_index
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip() == anchor_text:
            anchor = p
            break
    if anchor is None:
        raise ValueError("找不到锚点段落: %r" % anchor_text)
    el = tbl._tbl
    el.getparent().remove(el)      # 从原位置摘除
    anchor._p.addnext(el)          # 挂到锚点段落之后
    return tbl


def remove_table(doc, tbl_or_index):
    """按索引或对象整张删除表格（把 w:tbl 从 body 摘除）。用于清理重复旧表。"""
    tbl = doc.tables[tbl_or_index] if isinstance(tbl_or_index, int) else tbl_or_index
    el = tbl._tbl
    el.getparent().remove(el)
    return True


def insert_empty_paragraph_after(doc, element):
    """在给定 body 级元素后插入一个空段落，用来隔开相邻表格防 Word 自动合并。"""
    from docx.oxml import OxmlElement
    p = OxmlElement('w:p')
    element.addnext(p)
    return p


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "inspect"
    path = sys.argv[2]
    d = load(path)
    if cmd == "inspect":
        print("paragraphs=%d tables=%d" % (len(d.paragraphs), len(d.tables)))
        iter_body(d)
    elif cmd == "dump":
        idx = int(sys.argv[3]) if len(sys.argv) > 3 else 0
        print(read_table(d.tables[idx]))
